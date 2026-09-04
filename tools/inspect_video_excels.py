import sys
from openpyxl import load_workbook

PY = "/Users/jeremy/.workbuddy/binaries/python/envs/default/bin/python"

files = [
    "/Users/jeremy/Desktop/小学知识成长地图应用/苏e新课_小学一至四年级_知识点与课件资源汇总.xlsx",
    "/Users/jeremy/Desktop/小学知识成长地图应用/苏e优课_小学一至六年级_语文人教版_英语译林版_知识点与视频索引.xlsx",
]

for f in files:
    print("=" * 80)
    print("FILE:", f.split("/")[-1])
    wb = load_workbook(f, data_only=True, read_only=True)
    print("SHEETS:", wb.sheetnames)
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        print(f"\n--- sheet '{sn}' ({len(rows)} rows) ---")
        if not rows:
            continue
        for ri, r in enumerate(rows[:4]):
            cells = [("" if c is None else str(c))[:40] for c in r]
            print(f"  row{ri}:", " | ".join(cells))
    wb.close()
