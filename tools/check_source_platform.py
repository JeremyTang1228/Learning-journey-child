from openpyxl import load_workbook
import urllib.request, ssl

f2 = "/Users/jeremy/Desktop/小学知识成长地图应用/苏e优课_小学一至六年级_语文人教版_英语译林版_知识点与视频索引.xlsx"
wb = load_workbook(f2, data_only=True, read_only=True)
ws = wb["数据说明"]
print("=== 数据说明 ===")
for r in ws.iter_rows(values_only=True):
    print("  ", [("" if c is None else str(c))[:60] for c in r])
wb.close()

# 尝试用 chapter_id 拼源平台观看地址，看是否可达 & 能否内嵌
ctx = ssl.create_default_context(); ctx.check_hostname=False; ctx.verify_mode=ssl.CERT_NONE
candidates = [
    "https://basic.jiangsu.smartedu.cn/",
    "https://basic.jiangsu.smartedu.cn/#/clazz/courseDetail?chapterId=22",
    "https://basic.jiangsu.smartedu.cn/source?chapterId=22",
    "https://basic.jiangsu.smartedu.cn/tbkt/train/course/watch?chapterId=22",
]
print("\n=== 测试源平台 URL ===")
for u in candidates:
    try:
        req = urllib.request.Request(u, method="GET", headers={"User-Agent":"Mozilla/5.0"})
        resp = urllib.request.urlopen(req, timeout=12, context=ctx)
        body = resp.read(400).decode("utf-8","ignore")
        xfo = resp.headers.get("X-Frame-Options")
        csp = resp.headers.get("Content-Security-Policy")
        print(f"OK  {resp.status} xfo={xfo} csp={'yes' if csp else 'no'} {u[:70]}")
        print("     body:", body[:120].replace("\n"," "))
    except urllib.error.HTTPError as e:
        print(f"HTTP {e.code} {u[:70]}")
    except Exception as e:
        print(f"ERR {str(e)[:60]} {u[:70]}")
