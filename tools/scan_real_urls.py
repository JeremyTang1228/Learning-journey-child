from openpyxl import load_workbook
import urllib.request, ssl, re

f1 = "/Users/jeremy/Desktop/小学知识成长地图应用/苏e新课_小学一至四年级_知识点与课件资源汇总.xlsx"
f2 = "/Users/jeremy/Desktop/小学知识成长地图应用/苏e优课_小学一至六年级_语文人教版_英语译林版_知识点与视频索引.xlsx"

def scan(path, sheet, col):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    i = list(hdr).index(col)
    real = []
    placeholder = set()
    for r in rows[1:]:
        v = r[i]
        if v is None: continue
        s = str(v).strip()
        if s.startswith("http"):
            real.append(s)
        else:
            placeholder.add(s)
    print(f"{path.split('/')[-1]} | {sheet}.{col}: 真实URL={len(real)} 占位符种类={placeholder}")
    wb.close()
    return real

# 扫描所有视频相关链接列
for f in [f1, f2]:
    for sheet in ["视频索引", "知识点目录汇总"]:
        try:
            scan(f, sheet, "视频详情链接")
        except Exception as e:
            print("skip", sheet, e)
        try:
            scan(f, sheet, "视频链接")
        except Exception as e:
            print("skip", sheet, e)

# 测试：能否从 fileId 推导可播放视频地址？
# 取一个已知 fileId
fid = "fd959ad01397757894742704304"
base = "http://1258942858.vod2.myqcloud.com/5f629bc5vodtranssh1258942858"
candidates = [
    f"{base}/{fid}/{fid}.mp4",
    f"{base}/{fid}/v.f100.mp4",
    f"{base}/{fid}/video.mp4",
    f"{base}/{fid}/index.m3u8",
]
ctx = ssl.create_default_context()
ctx.check_hostname = False; ctx.verify_mode = ssl.CERT_NONE
print("\n=== 测试推导视频地址 ===")
for u in candidates:
    try:
        req = urllib.request.Request(u, method="HEAD", headers={"User-Agent":"Mozilla/5.0","Referer":"https://basic.jiangsu.smartedu.cn/"})
        resp = urllib.request.urlopen(req, timeout=12, context=ctx)
        print("OK ", resp.status, resp.headers.get("Content-Type"), u)
    except urllib.error.HTTPError as e:
        print("HTTP", e.code, u)
    except Exception as e:
        print("ERR", str(e)[:50], u)
