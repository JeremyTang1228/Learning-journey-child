#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Master Data V2 深度校验：ID 唯一性、外键完整性、状态枚举、排序字段、资源 URL 可用性。"""
from collections import Counter, defaultdict
from openpyxl import load_workbook

XLSX = "/Users/jeremy/Desktop/小学知识成长地图应用/02_Master_Data/小学知识成长地图_Master_Data_V2_内容管理版.xlsx"
wb = load_workbook(XLSX, data_only=True, read_only=True)


def sheet(name):
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    return [{header[i]: r[i] for i in range(min(len(header), len(r))) if header[i]} for r in rows[1:]]


subs = sheet("subject_config")
curs = sheet("curriculums")
units = sheet("units")
kns = sheet("knowledge_nodes")
ress = sheet("resources")

print("=" * 70)
print("A. 学科 / 教材版本 / 年级 分布")
print("  学科:", dict(Counter(r["subject"] for r in curs)))
print("  年级:", dict(Counter(r["grade"] for r in curs)))
print("  semester 取值:", dict(Counter(str(r["semester"]) for r in curs)))
print("  is_core_subject 非空:", sum(1 for r in curs if r.get("is_core_subject") not in (None, "")))
print("  curriculums.status:", dict(Counter(str(r["status"]) for r in curs)))

print("\n" + "=" * 70)
print("B. ID 唯一性")
for name, rows, key in (("curriculums", curs, "curriculum_id"), ("units", units, "unit_id"),
                        ("knowledge_nodes", kns, "knowledge_id"), ("resources", ress, "resource_id")):
    ids = [r[key] for r in rows]
    dup = [k for k, v in Counter(ids).items() if v > 1]
    print(f"  {name:<16} {len(ids):>5} 行, 重复 {len(dup)} 个  {dup[:3]}")

print("\n" + "=" * 70)
print("C. 外键完整性")
cid = {r["curriculum_id"] for r in curs}
uid = {r["unit_id"] for r in units}
kid = {r["knowledge_id"] for r in kns}
print("  units.curricular_id 缺失:", sum(1 for r in units if r["curriculum_id"] not in cid))
print("  kn.unit_id 缺失:", sum(1 for r in kns if r["unit_id"] not in uid))
print("  res.knowledge_id 孤儿:", sum(1 for r in ress if r["knowledge_id"] not in kid))
print("  从未被 resource 引用的 kn:", len(kid - {r['knowledge_id'] for r in ress}))

print("\n" + "=" * 70)
print("D. 状态枚举 (规范要求 draft/published/hidden/deprecated)")
for name, rows in (("curriculums", curs), ("units", units),
                   ("knowledge_nodes", kns), ("resources", ress), ("subject_config", subs)):
    print(f"  {name:<16}", dict(Counter(str(r.get("status")) for r in rows)))

print("\n" + "=" * 70)
print("E. 排序字段有效性 (sort_order / display_order / sequence)")
print("  units.sort_order 唯一值:", sorted(Counter(str(r.get("sort_order")) for r in units).items())[:10])
print("  kn.sort_order 唯一值:", sorted(Counter(str(r.get("sort_order")) for r in kns).items())[:10])
print("  kn.sequence 样例分布:", sorted(Counter(str(r.get("sequence")) for r in kns).items(), key=lambda x: x[0])[:12])
print("  subject.display_order:", dict(Counter(str(r.get("display_order")) for r in subs)))

# 同一 unit 下 sequence 是否重复
seq_dup = 0
by_unit = defaultdict(list)
for r in kns:
    by_unit[r["unit_id"]].append(str(r.get("sequence")))
for u, s in by_unit.items():
    if len(set(s)) != len(s):
        seq_dup += 1
print(f"  同一 unit 内 sequence 重复的 unit 数: {seq_dup} / {len(by_unit)}")

print("\n" + "=" * 70)
print("F. 资源 URL 可用性 (关键)")
urlc = Counter()
for r in ress:
    u = str(r.get("url") or "")
    if u.startswith("http"):
        urlc["http"] += 1
    elif u.strip() == "打开视频":
        urlc["占位:打开视频"] += 1
    elif u.strip() == "":
        urlc["空"] += 1
    else:
        urlc["其他:" + u[:12]] += 1
print("  ", dict(urlc))
print("  availability:", dict(Counter(str(r.get("availability")) for r in ress)))
print("  resource_type:", dict(Counter(str(r.get("resource_type")) for r in ress)))
print("  source_platform:", dict(Counter(str(r.get("source_platform")) for r in ress)))

print("\n" + "=" * 70)
print("G. 标题缺失")
print("  kn.title 空:", sum(1 for r in kns if not str(r.get("title") or "").strip()))
print("  kn data_quality:", dict(Counter(str(r.get("data_quality")) for r in kns)))
print("  kn data_origin:", dict(Counter(str(r.get("data_origin")) for r in kns)))

print("\n" + "=" * 70)
print("H. 每个知识点资源数分布")
per = Counter(r["knowledge_id"] for r in ress)
cnt = Counter(per.get(k, 0) for k in kid)
print("  ", dict(sorted(cnt.items())))

print("\n" + "=" * 70)
print("I. 一年级数学(苏教版上册) 链路抽样")
for c in curs:
    if c["grade"] == "一年级" and c["subject"] == "小学数学" and str(c["semester"]) == "上册":
        print("  curriculum:", c["curriculum_id"], c["textbook"], "map_name=", c.get("map_name"))
        us = [u for u in units if u["curriculum_id"] == c["curriculum_id"]][:4]
        for u in us:
            ks = [k for k in kns if k["unit_id"] == u["unit_id"]][:3]
            print(f"    Unit[{u['unit_name']}] id={u['unit_id']} sort={u.get('sort_order')} kn数={len(ks)}")
            for k in ks:
                n = per.get(k["knowledge_id"], 0)
                print(f"        - seq{k.get('sequence')} {k['title']} ({n}资源) {k['knowledge_id']}")
        break

print("\n" + "=" * 70)
print("J. 学科 -> 是否有 curriculum")
print("  subject_config 学科数:", len(subs))
for s in subs:
    n = sum(1 for c in curs if c["subject"] == s["subject"])
    print(f"    {s['subject']:<10} curriculum数={n:<3} order={s['display_order']} visible={s['is_visible']} status={s['status']}")
wb.close()
