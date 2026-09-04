from openpyxl import load_workbook

f1 = "/Users/jeremy/Desktop/小学知识成长地图应用/苏e新课_小学一至四年级_知识点与课件资源汇总.xlsx"
f2 = "/Users/jeremy/Desktop/小学知识成长地图应用/苏e优课_小学一至六年级_语文人教版_英语译林版_知识点与视频索引.xlsx"

def dump(path, sheet, cols, n=4):
    print("="*90)
    print("FILE:", path.split("/")[-1], "| SHEET:", sheet)
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    print("HEADERS:", list(hdr))
    idx = [list(hdr).index(c) for c in cols]
    for r in rows[1:n+1]:
        for i in idx:
            c = r[i]
            print(f"  [{hdr[i]}] = {c!r}")
        print("  " + "-"*40)
    wb.close()

# 苏e新课 视频索引
dump(f1, "视频索引", ["年级","学科","册次","单元 / 模块","课时 / 知识点","视频标题","视频详情链接","封面图"])
# 苏e优课 视频索引
dump(f2, "视频索引", ["年级","学科","册次","单元 / 模块","课时 / 知识点","视频标题","视频详情链接","封面图"])
# 苏e新课 知识点目录汇总 视频链接 + 封面图信息?
dump(f1, "知识点目录汇总", ["年级","学科","册次","单元 / 模块","课时 / 知识点","视频链接","视频详情链接" if "视频详情链接" in list(load_workbook(f1, read_only=True)["知识点目录汇总"].iter_rows(values_only=True))[0] else "本地文件夹"])
