#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""读取 Master Data V2 xlsx，输出每个 sheet 的列名、行数、样例数据与空值统计。"""
import sys
from openpyxl import load_workbook

XLSX = "/Users/jeremy/Desktop/小学知识成长地图应用/02_Master_Data/小学知识成长地图_Master_Data_V2_内容管理版.xlsx"

wb = load_workbook(XLSX, data_only=True, read_only=True)
print("=" * 78)
print("SHEETS:", wb.sheetnames)
print("=" * 78)

for ws in wb.worksheets:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print(f"\n### [{ws.title}]  (空表, 0 行)")
        continue
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    data = rows[1:]
    print(f"\n### [{ws.title}]  数据行数 = {len(data)}")
    print(f"列({len(header)}): {header}")

    # 每列的非空统计 + 样例值
    for i, col in enumerate(header):
        if not col:
            continue
        vals = [r[i] for r in data if i < len(r) and r[i] is not None and str(r[i]).strip() != ""]
        non_empty = len(vals)
        uniq = list(dict.fromkeys([str(v)[:28] for v in vals]))[:4]
        flag = "  ⚠️全空" if non_empty == 0 else ""
        print(f"   - {col:<28} 非空 {non_empty:>5}/{len(data):<5} 样例={uniq}{flag}")

    # 打印前 2 行数据
    if data:
        print("  前2行样例:")
        for r in data[:2]:
            cells = {header[i]: r[i] for i in range(min(len(header), len(r))) if header[i]}
            print("   ", str(cells)[:600])
wb.close()
